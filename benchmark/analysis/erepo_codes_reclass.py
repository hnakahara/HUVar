#!/usr/bin/env python3
"""Reviewer 3, Comment 15: eRepo variants whose own applied codes, combined with
HUVar's rules, do not reproduce the eRepo (VCEP) classification.

Each eRepo variant's "Applied Evidence Codes (Met)" (at their stated strengths;
an unsuffixed code = the 2015 default strength, i.e. unsuffixed PM2 = Moderate)
is classified with HUVar's own classifiers:
  - Bayesian point system (classifier_bayesian.py; HUVar's primary output)
  - ACMG/AMP 2015 Table 5, applied strictly (contradictory evidence -> VUS);
    HUVar's current classifier_2015.py result is listed alongside
Variants where the Bayesian class differs from the eRepo assertion are listed
with a cause category:
  A  benign and pathogenic codes together (points net out; 2015 rules -> VUS)
  B  one direction only; 2015 rules reproduce the eRepo class (point
     threshold differs from the 2015 combination table)
  C  neither HUVar rule set reproduces the eRepo class (VCEP-specific
     combining rule or expert judgement)

  python erepo_codes_reclass.py --dir <eRepo outputs with
      concordance_hg38_revel_manual_per_variant.tsv and
      concordance_hg38_revel_per_variant.tsv> --erepo erepo_classifications_hg38.tsv \
      --huvar-src <HUVar>/src --out <dir>
Writes erepo_codes_reclass_variants.tsv, erepo_codes_reclass_summary.txt and
Supplementary_Table_eRepo_codes_reclass.xlsx.
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from pathlib import Path

import pandas as pd

CLS = ["P", "LP", "VUS", "LB", "B"]
RANK = {"P": 2, "LP": 1, "VUS": 0, "LB": -1, "B": -2}
ST = {"": None, "very": "VERY_STRONG", "verystrong": "VERY_STRONG", "very strong": "VERY_STRONG",
      "very_strong": "VERY_STRONG", "strong": "STRONG", "moderate": "MODERATE",
      "supporting": "SUPPORTING", "stand alone": "VERY_STRONG", "standalone": "VERY_STRONG",
      "stand": "VERY_STRONG", "stand_alone": "VERY_STRONG"}
CAUSE = {"A": "A: benign + pathogenic codes combined",
         "B": "B: ACMG/AMP 2015 Table 5 reproduces eRepo (point threshold differs)",
         "C": "C: neither rule set reproduces eRepo (VCEP-specific rule / judgement)"}


def table5(res) -> str:
    """Richards et al. 2015 Table 5, applied strictly (strength-aware buckets;
    benign Moderate/ThreePoint counted as Supporting, as in classifier_2015).
    Contradictory benign and pathogenic criteria -> VUS."""
    b = Counter()
    for r in res:
        s, path = r.strength.name, r.direction.name == "PATHOGENIC"
        if path:
            k = {"VERY_STRONG": "pvs", "STRONG": "ps", "THREE_POINT": "pm", "MODERATE": "pm",
                 "SUPPORTING": "pp"}.get(s)
        else:
            k = {"VERY_STRONG": "ba", "STRONG": "bs"}.get(s, "bp")
        if k:
            b[k] += 1
    pvs, ps, pm, pp, ba, bs, bp = (b[k] for k in ("pvs", "ps", "pm", "pp", "ba", "bs", "bp"))
    P = ((pvs >= 1 and (ps >= 1 or pm >= 2 or (pm >= 1 and pp >= 1) or pp >= 2)) or ps >= 2
         or (ps >= 1 and (pm >= 3 or (pm >= 2 and pp >= 2) or (pm >= 1 and pp >= 4))))
    LP = ((pvs >= 1 and pm == 1) or (ps == 1 and 1 <= pm <= 2) or (ps == 1 and pp >= 2) or pm >= 3
          or (pm == 2 and pp >= 2) or (pm == 1 and pp >= 4))
    B = ba >= 1 or bs >= 2
    LB = (bs == 1 and bp >= 1) or bp >= 2
    if (P or LP) and (B or LB):
        return "VUS"
    return "P" if P else "LP" if LP else "B" if B else "LB" if LB else "VUS"


def cat(t: str, p: str) -> str:
    if t == p:
        return "exact"
    a, b = RANK[t], RANK[p]
    return "confidence" if a * b > 0 else ("significant" if a == 0 or b == 0 else "opposite")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", type=Path, required=True)
    ap.add_argument("--erepo", type=Path, required=True)
    ap.add_argument("--huvar-src", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--results-dir", type=Path, default=None,
                    help="folder with erepo_hg38_huhvar_revel[_manual]_results.tsv (adds HUVar's criteria)")
    ap.add_argument("--build", default="hg38", choices=["hg38", "hg19"],
                    help="assembly of the input files (concordance_<build>_*, erepo_<build>_huvar_*)")
    a = ap.parse_args()
    sys.path.insert(0, str(a.huvar_src))
    from acmg_classifier.classification.classifier_2015 import Classifier2015
    from acmg_classifier.classification.classifier_bayesian import ClassifierBayesian
    from acmg_classifier.models.criteria import CriteriaResult
    from acmg_classifier.models.enums import ACMGCriterion, CriterionStrength, Pathogenicity

    PATH = {Pathogenicity.PATHOGENIC: "P", Pathogenicity.LIKELY_PATHOGENIC: "LP",
            Pathogenicity.VUS: "VUS", Pathogenicity.LIKELY_BENIGN: "LB", Pathogenicity.BENIGN: "B"}
    bay, c15 = ClassifierBayesian(), Classifier2015()
    unparsed = Counter()

    def results(codes: str):
        out, labels = [], []
        for t in str(codes).split(","):
            t = t.strip()
            if not t or t == "nan":
                continue
            m = re.match(r"([A-Z]+\d)", t)
            if not m:
                unparsed[t] += 1
                continue
            base = m.group(1)
            sfx = t[len(base):].lstrip("_ ").lower()
            if sfx not in ST:
                unparsed[t] += 1
                continue
            try:
                crit = ACMGCriterion[base]
            except KeyError:
                unparsed[t] += 1
                continue
            st = ST[sfx]
            if st is None and base == "PM2":
                st = "MODERATE"  # eRepo unsuffixed PM2 = 2015 Moderate
            strength = CriterionStrength[st] if st else None
            r = CriteriaResult.met(crit, strength)
            out.append(r)
            labels.append(f"{t}({r.points:+d})")
        return out, " ".join(labels)

    pv = pd.read_csv(a.dir / f"concordance_{a.build}_revel_manual_per_variant.tsv", sep="\t", dtype=str)
    pv = pv[pv.truth.isin(CLS)].drop_duplicates("variant_id")
    auto = pd.read_csv(a.dir / f"concordance_{a.build}_revel_per_variant.tsv", sep="\t", dtype=str)
    auto = auto.drop_duplicates("variant_id").set_index("variant_id").pred
    er = pd.read_csv(a.erepo, sep="\t", comment="#", dtype=str).drop_duplicates("variant_id").set_index("variant_id")

    # HUVar's own criteria (met, with strength) for each configuration
    SHORTS = {"VeryStrong": "VS", "Strong": "S", "ThreePoint": "3pt", "Moderate": "M", "Supporting": "P",
              "StandAlone": "SA"}
    PTS = {"VeryStrong": 8, "Strong": 4, "ThreePoint": 3, "Moderate": 2, "Supporting": 1, "StandAlone": 8}
    CODES = ["PVS1", "PS1", "PS2", "PS3", "PS4", "PM1", "PM2", "PM3", "PM4", "PM5", "PM6", "PP1", "PP2",
             "PP3", "PP4", "BA1", "BS1", "BS2", "BS3", "BS4", "BP1", "BP2", "BP3", "BP4", "BP5", "BP7"]
    hu = {}
    rdir = a.results_dir or a.dir
    for cfg in ("revel_manual", "revel"):
        f = rdir / f"erepo_{a.build}_huhvar_{cfg}_results.tsv"
        if not f.exists():
            print(f"note: {f} not found - HUVar criteria for {cfg} omitted", file=sys.stderr)
            continue
        t = pd.read_csv(f, sep="\t", dtype=str, comment="#",
                        usecols=lambda c: c in ["variant_id", "bayesian_score"] + CODES + [c_ + "_strength" for c_ in CODES])
        t = t.drop_duplicates("variant_id").set_index("variant_id")
        m = {}
        for vid, row in t.iterrows():
            met = {}
            for c in CODES:
                if str(row.get(c)) == "1":
                    met[c] = str(row.get(c + "_strength", ""))
            m[vid] = (met, row.get("bayesian_score", ""))
        hu[cfg] = m

    def lab_hu(met: dict) -> str:
        return ",".join(f"{c}_{SHORTS.get(s, s)}({'-' if c.startswith('B') else '+'}{PTS.get(s, 0)})"
                        for c, s in met.items())

    def norm_er(res) -> dict:
        return {x.criterion.name: x.strength.name for x in res}

    HUST = {"VeryStrong": "VERY_STRONG", "Strong": "STRONG", "ThreePoint": "THREE_POINT",
            "Moderate": "MODERATE", "Supporting": "SUPPORTING", "StandAlone": "VERY_STRONG"}

    def diff(erd: dict, met: dict) -> tuple[str, str, str]:
        h = {c: HUST.get(s, s) for c, s in met.items()}
        miss = [c for c in erd if c not in h]
        extra = [c for c in h if c not in erd]
        st = [f"{c} {erd[c].lower()}->{h[c].lower()}" for c in erd if c in h and h[c] != erd[c]
              and c != "BA1"]
        return ",".join(miss), ",".join(extra), "; ".join(st)

    rows = []
    for r in pv.itertuples(index=False):
        codes = getattr(r, "_8")  # "Applied Evidence Codes (Met)"
        res, lab = results(codes)
        score, b = bay.classify(res)
        h15, _ = c15.classify(res)
        b, h15, k15 = PATH[b], PATH[h15], table5(res)
        rec = {"variant_id": r.variant_id, "gene": r.Gene,
               "hgvs_c": er.hgvs_c.get(r.variant_id, ""), "expert_panel": er.expert_panel.get(r.variant_id, ""),
               "disease": er.disease.get(r.variant_id, ""), "approval_date": er.approval_date.get(r.variant_id, ""),
               "eRepo": r.truth, "eRepo_codes": codes, "codes_points": lab, "score": score,
               "codes_bayesian": b, "codes_2015": k15, "codes_huvar2015": h15, "difference": cat(r.truth, b),
               "HUVar_revel_manual": r.pred, "HUVar_revel_auto": auto.get(r.variant_id, "")}
        erd = norm_er(res)
        for cfg, tag in (("revel_manual", "curated"), ("revel", "auto")):
            if cfg in hu and r.variant_id in hu[cfg]:
                met, sc = hu[cfg][r.variant_id]
                rec[f"HUVar_{tag}_criteria"] = lab_hu(met)
                rec[f"HUVar_{tag}_score"] = sc
                rec[f"{tag}_missing_vs_eRepo"], rec[f"{tag}_extra_vs_eRepo"], rec[f"{tag}_strength_diff"] = diff(erd, met)
        if b != r.truth:
            dirs = {x.direction.name for x in res}
            if len(dirs) > 1:
                rec["cause"] = "A"
            elif k15 == r.truth:
                rec["cause"] = "B"
            else:
                rec["cause"] = "C"
        rows.append(rec)
    d = pd.DataFrame(rows)
    x = d[d.codes_bayesian != d.eRepo].copy()
    x["direction"] = ["HUVar higher" if RANK[p] > RANK[t] else "HUVar lower" for t, p in zip(x.eRepo, x.codes_bayesian)]
    x = x.sort_values(["cause", "eRepo", "codes_bayesian", "gene"])
    a.out.mkdir(parents=True, exist_ok=True)
    x.to_csv(a.out / "erepo_codes_reclass_variants.tsv", sep="\t", index=False)

    L = []
    p = L.append
    n = len(d)
    p(f"eRepo variants: {n}; codes -> HUVar Bayesian = eRepo: {(d.codes_bayesian == d.eRepo).sum()} "
      f"({100 * (d.codes_bayesian == d.eRepo).mean():.1f}%); 2015 rules = eRepo: "
      f"{(d.codes_2015 == d.eRepo).sum()} ({100 * (d.codes_2015 == d.eRepo).mean():.1f}%); HUVar's current "
      f"2015 classifier = eRepo: {(d.codes_huvar2015 == d.eRepo).sum()} ({100 * (d.codes_huvar2015 == d.eRepo).mean():.1f}%)")
    dd = d[d.codes_huvar2015 != d.codes_2015]
    p(f"HUVar 2015 classifier differs from strict Table 5 on the eRepo codes: {len(dd)} variants; "
      f"{dd.groupby(['codes_2015', 'codes_huvar2015']).size().to_dict()} (Table 5 -> HUVar); of these, "
      f"Table 5 = eRepo {(dd.codes_2015 == dd.eRepo).sum()}, HUVar 2015 = eRepo {(dd.codes_huvar2015 == dd.eRepo).sum()}")
    p(f"Bayesian class differs from eRepo: {len(x)} ({100 * len(x) / n:.1f}%) — "
      f"{x.difference.value_counts().to_dict()}")
    p("by cause: " + "; ".join(f"{CAUSE[k]} {v}" for k, v in x.cause.value_counts().sort_index().items()))
    p("")
    p("eRepo -> class from eRepo codes (Bayesian), by cause:")
    t = x.groupby(["eRepo", "codes_bayesian", "cause"]).size().unstack(fill_value=0)
    t["total"] = t.sum(axis=1)
    p(t.sort_values("total", ascending=False).to_string())
    p("")
    p("most frequent code combinations (n >= 5):")
    combos = x.groupby(["eRepo", "codes_bayesian", "cause", "eRepo_codes"]).size().sort_values(ascending=False)
    for (t_, b_, c_, cc), v in combos[combos >= 5].items():
        p(f"  {v:4d}  eRepo {t_:3s} -> {b_:3s} [{c_}]  {cc}")
    p("")
    p("by expert panel (n differing / n in eRepo):")
    ep_tot = d.assign(ep=d.expert_panel.fillna("")).groupby("ep").size()
    ep_x = x.assign(ep=x.expert_panel.fillna("")).groupby("ep").size().sort_values(ascending=False)
    for ep, v in ep_x.items():
        p(f"  {v:4d} / {ep_tot[ep]:5d} ({100 * v / ep_tot[ep]:5.1f}%)  {ep}")
    p("")
    p("HUVar (REVEL + curated) result for these variants: "
      f"same as eRepo {(x.HUVar_revel_manual == x.eRepo).sum()}, same as the codes-derived class "
      f"{(x.HUVar_revel_manual == x.codes_bayesian).sum()}, other {((x.HUVar_revel_manual != x.eRepo) & (x.HUVar_revel_manual != x.codes_bayesian)).sum()}")
    for tag in ("curated", "auto"):
        if f"{tag}_missing_vs_eRepo" not in x:
            continue
        same = ((x[f"{tag}_missing_vs_eRepo"].fillna("") == "") & (x[f"{tag}_extra_vs_eRepo"].fillna("") == "")
                & (x[f"{tag}_strength_diff"].fillna("") == ""))
        cnt = lambda col: Counter(c for v in x[col].fillna("") for c in str(v).split(",") if c)
        stc = Counter(v.split(" ")[0] for s_ in x[f"{tag}_strength_diff"].fillna("") for v in s_.split("; ") if v)
        p(f"HUVar {tag} criteria identical to the eRepo codes (code and strength): {same.sum()} / {len(x)}; "
          f"eRepo code absent in HUVar {dict(cnt(f'{tag}_missing_vs_eRepo').most_common(8))}; "
          f"HUVar-only {dict(cnt(f'{tag}_extra_vs_eRepo').most_common(8))}; strength differs {dict(stc.most_common(6))}")
    if unparsed:
        p(f"unparsed code tokens: {dict(unparsed.most_common(20))}")
    txt = "\n".join(L)
    (a.out / "erepo_codes_reclass_summary.txt").write_text(txt + "\n", encoding="utf-8")
    print(txt)

    # xlsx
    from openpyxl import Workbook
    from openpyxl.styles import Font
    bold = Font(bold=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["eRepo variants whose own applied codes, combined with HUVar's Bayesian point system, "
               "give a different classification from the eRepo assertion (GRCh38)."])
    ws["A1"].font = bold
    ws.append([f"n = {len(x)} of {n} ({100 * len(x) / n:.1f}%). Codes taken at their stated strengths; "
               "unsuffixed codes at the ACMG/AMP 2015 default (PM2 = Moderate)."])
    ws.append([])
    ws.append(["Cause", "Description", "n"])
    for c in "ABC":
        ws.append([c, CAUSE[c][3:], int((x.cause == c).sum())])
    ws.append([])
    ws.append(["eRepo", "Class from eRepo codes (Bayesian)", "A", "B", "C", "Total"])
    for c in ws[ws.max_row]:
        c.font = bold
    tt = x.groupby(["eRepo", "codes_bayesian", "cause"]).size().unstack(fill_value=0).reindex(columns=list("ABC"), fill_value=0)
    tt["Total"] = tt.sum(axis=1)
    for (e, b), r in tt.sort_values("Total", ascending=False).iterrows():
        ws.append([e, b] + [int(v) for v in r])
    ws.append([])
    ws.append(["Expert panel", "n differing", "n in eRepo", "%"])
    for c in ws[ws.max_row]:
        c.font = bold
    for ep, v in ep_x.items():
        ws.append([ep, int(v), int(ep_tot[ep]), round(100 * v / ep_tot[ep], 1)])
    for c in ws[4]:
        c.font = bold
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 40

    w2 = wb.create_sheet("Code combinations")
    w2.append(["n", "eRepo", "Class from codes", "Cause", "eRepo applied codes"])
    for (t_, b_, c_, cc), v in combos.items():
        w2.append([int(v), t_, b_, c_, cc])
    w3 = wb.create_sheet("Variants")
    cols = ["variant_id", "gene", "hgvs_c", "expert_panel", "approval_date", "eRepo", "eRepo_codes",
            "codes_points", "score", "codes_bayesian", "codes_2015", "codes_huvar2015", "difference", "direction", "cause",
            "HUVar_revel_manual", "HUVar_curated_score", "HUVar_curated_criteria", "curated_missing_vs_eRepo",
            "curated_extra_vs_eRepo", "curated_strength_diff",
            "HUVar_revel_auto", "HUVar_auto_score", "HUVar_auto_criteria", "auto_missing_vs_eRepo",
            "auto_extra_vs_eRepo", "auto_strength_diff"]
    cols = [c for c in cols if c in x.columns]
    hdr = ["variant_id", "gene", "HGVS c.", "expert panel", "approval date", "eRepo class", "eRepo applied codes",
           "codes (points)", "score", "class from codes (Bayesian)", "class from codes (ACMG 2015 Table 5)",
           "class from codes (HUVar 2015 classifier, current)",
           "difference", "direction", "cause",
           "HUVar REVEL + curated: class", "score", "criteria (points)", "eRepo code not in HUVar",
           "HUVar code not in eRepo", "strength differs (eRepo->HUVar)",
           "HUVar REVEL automated: class", "score", "criteria (points)", "eRepo code not in HUVar",
           "HUVar code not in eRepo", "strength differs (eRepo->HUVar)"]
    hdr = [h for h, c in zip(hdr, ["variant_id", "gene", "hgvs_c", "expert_panel", "approval_date", "eRepo",
           "eRepo_codes", "codes_points", "score", "codes_bayesian", "codes_2015", "codes_huvar2015", "difference",
           "direction", "cause", "HUVar_revel_manual", "HUVar_curated_score", "HUVar_curated_criteria",
           "curated_missing_vs_eRepo", "curated_extra_vs_eRepo", "curated_strength_diff", "HUVar_revel_auto",
           "HUVar_auto_score", "HUVar_auto_criteria", "auto_missing_vs_eRepo", "auto_extra_vs_eRepo",
           "auto_strength_diff"]) if c in cols]
    w3.append(hdr)
    for r in x[cols].itertuples(index=False):
        w3.append([None if pd.isna(v) else v for v in r])
    for w in (w2, w3):
        for c in w[1]:
            c.font = bold
        w.freeze_panes = "A2"
    wb.save(a.out / "Supplementary_Table_eRepo_codes_reclass.xlsx")


if __name__ == "__main__":
    main()
