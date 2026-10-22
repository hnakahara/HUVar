#!/usr/bin/env python
"""Format analysis_criteria_metrics.tsv into a combined table (Excel): 6 methods x implemented codes x 4 metrics.
Cells are merged by build (hg19/hg38) and method; only implemented codes are listed."""
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path("outputs/analysis_criteria_metrics.tsv")
OUT = Path("outputs/criteria_metrics_table.xlsx")

METHOD_ORDER = [
    "InterVar",
    "BIAS-2015 (default)",
    "BIAS-2015 (no PP5/BP6)",
    "HUVar ESM1b+OpenSpliceAI",
    "HUVar AlphaMissense+SpliceAI",
    "HUVar REVEL/BayesDel/CADD+SpliceAI",
]
CODE_ORDER = ["PVS1", "PS1", "PS2", "PS3", "PS4",
              "PM1", "PM2", "PM3", "PM4", "PM5", "PM6",
              "PP1", "PP2", "PP3", "PP4", "PP5",
              "BA1", "BS1", "BS2", "BS3", "BS4",
              "BP1", "BP2", "BP3", "BP4", "BP5", "BP6", "BP7"]
BUILD_ORDER = ["hg19", "hg38"]
METRICS = [("sensitivity", "Sensitivity"), ("specificity", "Specificity"),
           ("precision", "Precision"), ("f1", "F1 Score")]

# data: (build, method) -> {code: row}
data = defaultdict(dict)
with SRC.open(encoding="utf-8") as f:
    for r in csv.DictReader(f, delimiter="\t"):
        data[(r["build"], r["method"])][r["code"]] = r

thin = Side(style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor="305496")
hdr_font = Font(bold=True, color="FFFFFF")
build_fill = PatternFill("solid", fgColor="D9E1F2")
name_fill = PatternFill("solid", fgColor="F2F2F2")
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center")

wb = Workbook()
ws = wb.active
ws.title = "CriteriaMetrics"

headers = ["Build", "Classifier", "Code"] + [m[1] for m in METRICS]
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(1, c)
    cell.fill, cell.font, cell.alignment, cell.border = hdr_fill, hdr_font, center, border

row_idx = 2
for build in BUILD_ORDER:
    build_start = row_idx
    for method in METHOD_ORDER:
        codes = data.get((build, method), {})
        ordered = [c for c in CODE_ORDER if c in codes]
        if not ordered:
            continue
        method_start = row_idx
        for code in ordered:
            r = codes[code]
            ws.cell(row_idx, 1, build).alignment = center
            ws.cell(row_idx, 2, method).alignment = left
            ws.cell(row_idx, 3, code).alignment = center
            for j, (key, _) in enumerate(METRICS, start=4):
                v = r[key].strip()
                cell = ws.cell(row_idx, j, float(v) if v else "N/A")
                cell.number_format = "0.000"
                cell.alignment = center
            for c in range(1, len(headers) + 1):
                ws.cell(row_idx, c).border = border
            row_idx += 1
        ws.merge_cells(start_row=method_start, start_column=2,
                       end_row=row_idx - 1, end_column=2)
        mc = ws.cell(method_start, 2)
        mc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        mc.fill = name_fill
    ws.merge_cells(start_row=build_start, start_column=1,
                   end_row=row_idx - 1, end_column=1)
    bc = ws.cell(build_start, 1)
    bc.alignment = center
    bc.fill = build_fill
    bc.font = Font(bold=True)

ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 8
for c in range(4, len(headers) + 1):
    ws.column_dimensions[get_column_letter(c)].width = 13
ws.freeze_panes = "A2"

wb.save(OUT)
print(f"Wrote {OUT}  ({row_idx - 2} data rows)")
