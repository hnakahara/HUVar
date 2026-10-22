#!/usr/bin/env python
"""Format analysis_category_metrics.tsv into a combined table (Excel): 9 classifiers x 3 categories x 4 metrics."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path("outputs/analysis_category_metrics.tsv")
OUT = Path("outputs/classifier_metrics_table.xlsx")

# method name in the metrics file -> display name (order of the 9 classifiers)
METHOD_ORDER = [
    ("InterVar",                 "InterVar"),
    ("BIAS-2015",                "BIAS-2015 (default)"),
    ("BIAS-2015 (no PP5/BP6)",   "BIAS-2015 (no PP5/BP6)"),
    ("HUHVar default (2015)",    "HUVar ESM1b+OpenSpliceAI (ACMG 2015)"),
    ("HUHVar default (bayesian)","HUVar ESM1b+OpenSpliceAI (Bayesian)"),
    ("HUHVar spliceAI (2015)",   "HUVar AlphaMissense+SpliceAI (ACMG 2015)"),
    ("HUHVar spliceAI (bayesian)","HUVar AlphaMissense+SpliceAI (Bayesian)"),
    ("HUHVar REVEL (2015)",      "HUVar REVEL/BayesDel/CADD+SpliceAI (ACMG 2015)"),
    ("HUHVar REVEL (bayesian)",  "HUVar REVEL/BayesDel/CADD+SpliceAI (Bayesian)"),
]
CATEGORY_ORDER = ["P/LP", "VUS", "B/LB"]
BUILD_ORDER = ["hg19", "hg38"]
METRICS = [("sensitivity", "Sensitivity"),
           ("specificity", "Specificity"),
           ("precision", "Precision"),
           ("f1", "F1 Score")]

# --- read data: (build, method, category) -> row dict ---
data = {}
with SRC.open(encoding="utf-8") as f:
    for r in csv.DictReader(f, delimiter="\t"):
        data[(r["build"], r["method"], r["category"])] = r

# --- styles ---
thin = Side(style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor="305496")
hdr_font = Font(bold=True, color="FFFFFF")
build_fill = PatternFill("solid", fgColor="D9E1F2")
name_fill = PatternFill("solid", fgColor="F2F2F2")
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center")

wb = Workbook()
ws = wb.active
ws.title = "Metrics"

# header row
headers = ["Build", "Classifier", "Category"] + [m[1] for m in METRICS]
ws.append(headers)
for c, _ in enumerate(headers, 1):
    cell = ws.cell(1, c)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = center
    cell.border = border

row_idx = 2
for build in BUILD_ORDER:
    build_start = row_idx
    for src_method, disp_method in METHOD_ORDER:
        method_start = row_idx
        for cat in CATEGORY_ORDER:
            r = data.get((build, src_method, cat))
            ws.cell(row_idx, 1, build).alignment = center
            ws.cell(row_idx, 2, disp_method).alignment = left
            ws.cell(row_idx, 3, cat).alignment = center
            for j, (key, _) in enumerate(METRICS, start=4):
                cell = ws.cell(row_idx, j, float(r[key]) if r else None)
                cell.number_format = "0.000"
                cell.alignment = center
            for c in range(1, len(headers) + 1):
                ws.cell(row_idx, c).border = border
            row_idx += 1
        # merge the classifier-name cells vertically
        ws.merge_cells(start_row=method_start, start_column=2,
                       end_row=row_idx - 1, end_column=2)
        ws.cell(method_start, 2).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True)
        ws.cell(method_start, 2).fill = name_fill
    # merge the build cells vertically
    ws.merge_cells(start_row=build_start, start_column=1,
                   end_row=row_idx - 1, end_column=1)
    ws.cell(build_start, 1).alignment = center
    ws.cell(build_start, 1).fill = build_fill
    ws.cell(build_start, 1).font = Font(bold=True)

# column widths
ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 46
ws.column_dimensions["C"].width = 9
for c in range(4, len(headers) + 1):
    ws.column_dimensions[get_column_letter(c)].width = 13

ws.freeze_panes = "A2"

wb.save(OUT)
print(f"Wrote {OUT}  ({len(BUILD_ORDER)} builds x {len(METHOD_ORDER)} classifiers x {len(CATEGORY_ORDER)} categories)")
